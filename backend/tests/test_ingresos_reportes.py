import datetime as dt
import io
import shutil
import subprocess
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Pago

D = dt.timedelta


@pytest.fixture
def datos(db, crear_usuario, fabrica, hoy_fijo):
    """Escenario fijo. Hoy = 20/10/2026 (2da quincena en curso)."""
    hoy_fijo(dt.date(2026, 10, 20))
    ana, beto = crear_usuario("ana.ruiz"), crear_usuario("beto.luna")
    a, b = fabrica.cliente(ana, "Cliente A"), fabrica.cliente(beto, "Cliente B")
    c = fabrica.cliente(None, "Cliente Huérfano")
    oct_ = lambda d: dt.date(2026, 10, d)  # noqa: E731
    e = {}
    e["a1"] = fabrica.ciclo(a, oct_(5), costo=1000, estado="renovado", decision="si", pagos=[(1000, oct_(2))])
    e["a2"] = fabrica.ciclo(a, oct_(18), costo=2000, estado="vencido", pagos=[(500, oct_(10))],
                            prorroga=(oct_(15), oct_(19)))                      # prórroga ya vencida
    e["a3"] = fabrica.ciclo(a, oct_(25), costo=3000)
    e["a4"] = fabrica.ciclo(a, dt.date(2026, 9, 20), costo=500, estado="vencido", pagos=[(100, dt.date(2026, 9, 1))],
                            prorroga=(dt.date(2026, 10, 10), oct_(25)))          # de septiembre, prórroga activa
    e["a5"] = fabrica.ciclo(a, oct_(10), costo=700, estado="archivado", decision="no")
    e["b1"] = fabrica.ciclo(b, oct_(30), costo=1500, pagos=[(1500, oct_(19))])
    e["c1"] = fabrica.ciclo(c, oct_(12), costo=800, por=ana)
    return dict(ana=ana, beto=beto, e=e)


def test_proyeccion_cobrado_y_pendiente_del_periodo_en_curso(api, auth, crear_usuario, datos):
    adm = auth(crear_usuario("admin.demo", rol="admin"))
    r = api.get("/api/ingresos?anio=2026&mes=10&quincena=ambas", headers=adm).json()
    t = r["totales"]
    assert (t["proyeccion"], t["cobrado"], t["pendiente"], t["pct_cobrado"]) == (8300, 3000, 5300, 36.1)
    assert r["periodo"]["cerrado"] is False and r["periodo"]["estado_texto"] == "Periodo parcial, corte al 20/10/2026"
    assert r["total_general"]["proyeccion"] == 8300
    por = {c["cm"]: c for c in r["por_cm"]}
    assert (por["Ana"]["proyeccion"], por["Ana"]["cobrado"]) == (6000, 1500)   # a5 archivado y a4 (sept) fuera
    assert (por["Beto"]["proyeccion"], por["Beto"]["cobrado"]) == (1500, 1500) and por["Beto"]["pct_cobrado"] == 100
    assert por["Por reasignar"]["proyeccion"] == 800


@pytest.mark.parametrize("q,proy,cobrado", [("1", 1800, 1000), ("2", 6500, 2000)])
def test_filtro_por_quincena(api, auth, crear_usuario, datos, q, proy, cobrado):
    adm = auth(crear_usuario("admin.demo", rol="admin"))
    t = api.get(f"/api/ingresos?anio=2026&mes=10&quincena={q}", headers=adm).json()["totales"]
    assert (t["proyeccion"], t["cobrado"]) == (proy, cobrado)


def test_periodo_cerrado_dice_cerrado_y_septiembre_solo_trae_sus_paquetes(api, auth, crear_usuario, datos, hoy_fijo):
    hoy_fijo(dt.date(2026, 11, 5))
    adm = auth(crear_usuario("admin.demo", rol="admin"))
    oct_ = api.get("/api/ingresos?anio=2026&mes=10&quincena=ambas", headers=adm).json()
    assert oct_["periodo"]["cerrado"] is True and oct_["periodo"]["estado_texto"] == "Periodo cerrado"
    sep = api.get("/api/ingresos?anio=2026&mes=9&quincena=2", headers=adm).json()
    assert sep["totales"]["proyeccion"] == 500 and sep["totales"]["cobrado"] == 100


def test_cobrado_solo_cuenta_hasta_la_fecha_de_corte(api, db, auth, crear_usuario, datos):
    """Un pago fechado después del corte (no debería existir, pero si existe) no entra a lo cobrado."""
    p = datos["e"]["a3"]
    db.add(Pago(paquete_id=p.id, monto=Decimal(999), fecha=dt.date(2026, 10, 21), registrado_por=datos["ana"].id))
    db.flush()
    db.refresh(p)
    t = api.get("/api/ingresos?anio=2026&mes=10&quincena=ambas", headers=auth(crear_usuario("admin.demo", rol="admin"))).json()
    assert t["totales"]["cobrado"] == 3000


def test_ingresos_por_cliente_suma_sus_paquetes(api, auth, datos):
    r = api.get("/api/ingresos?anio=2026&mes=10&quincena=ambas", headers=auth(datos["ana"])).json()
    cli = {c["cliente"]: c for c in r["clientes"]}
    assert list(cli) == ["Cliente A"] and cli["Cliente A"]["paquetes"] == 3
    assert (cli["Cliente A"]["proyeccion"], cli["Cliente A"]["actual"]) == (6000, 1500)
    assert "por_cm" not in r                                      # el CM no ve el desglose por CM


def test_cm_solo_ve_su_cartera_en_ingresos_y_reportes_aunque_pida_otro_cm(api, auth, datos):
    h = auth(datos["beto"])
    r = api.get("/api/ingresos?anio=2026&mes=10&quincena=ambas&cm_id=" + str(datos["ana"].id), headers=h).json()
    assert r["totales"]["proyeccion"] == 1500 and [c["cliente"] for c in r["clientes"]] == ["Cliente B"]
    rep = api.get("/api/reportes/datos?anio=2026&mes=10&quincena=ambas&cm_id=" + str(datos["ana"].id), headers=h).json()
    assert {f["cliente"] for f in rep["detalle"]} == {"Cliente B"}
    assert "Cliente A" not in api.get("/api/reportes/datos?anio=2026&mes=10&quincena=ambas", headers=h).text
    assert api.get("/api/reportes/excel?anio=2026&mes=10&quincena=ambas", headers=h).content.count(b"Cliente A") == 0


def test_admin_puede_filtrar_por_cm(api, auth, crear_usuario, datos):
    adm = auth(crear_usuario("admin.demo", rol="admin"))
    r = api.get(f"/api/ingresos?anio=2026&mes=10&quincena=ambas&cm_id={datos['beto'].id}", headers=adm).json()
    assert r["totales"]["proyeccion"] == 1500


def test_tasa_de_renovacion_renovados_entre_vencidos(api, auth, crear_usuario, datos):
    r = api.get("/api/ingresos?anio=2026&mes=10&quincena=ambas", headers=auth(crear_usuario("admin.demo", rol="admin"))).json()
    tasa = {t["cm"]: t for t in r["tasa_renovacion"]}
    # Ana: a1 (renovado), a2 (vencido), a5 (archivado) llegaron a su fecha; a3 (25 oct) aún no
    assert (tasa["Ana"]["vencidos"], tasa["Ana"]["renovados"], tasa["Ana"]["tasa"]) == (3, 1, 33.3)
    assert (tasa["Beto"]["vencidos"], tasa["Beto"]["renovados"]) == (0, 0)       # b1 vence el 30
    assert tasa["Por reasignar"]["vencidos"] == 1                                            # c1 llegó a su fecha
    assert r["tasa_total"] == {"vencidos": 4, "renovados": 1, "tasa": 25.0}


def test_reporte_datos_secciones(api, auth, crear_usuario, datos):
    d = api.get("/api/reportes/datos?anio=2026&mes=10&quincena=ambas", headers=auth(crear_usuario("admin.demo", rol="admin"))).json()
    assert set(d) >= {"periodo", "resumen", "por_cm", "detalle", "prorrogas", "pendientes", "tasa_renovacion"}
    assert len(d["detalle"]) == 5                    # 7 ciclos - a4 (septiembre) - a5 (archivado)
    # prórrogas: NO se limitan al periodo (a4 es de septiembre); vencidas primero
    pr = d["prorrogas"]
    assert [(p["paquete"], p["vencida"], p["dias"]) for p in pr] == [("Básico", True, -1), ("Básico", False, 5)]
    assert [p["cliente"] for p in pr] == ["Cliente A", "Cliente A"] and float(pr[0]["restante"]) == 1500
    # pendientes agrupados por CM: a2 vencido sin decisión; a4 también (septiembre, arrastrado)
    venc = d["pendientes"]["vencidos"]
    assert [g["cm"] for g in venc] == ["Ana"] and len(venc[0]["items"]) == 2
    assert d["pendientes"]["renovados_sin_pago"] == []
    assert d["pendientes"]["por_vencer"] == []                                  # a3 vence el 25: faltan 5 días (> 4)


def test_endpoints_validan_periodo(api, auth, datos):
    h = auth(datos["ana"])
    assert api.get("/api/ingresos?mes=13", headers=h).status_code == 422
    assert api.get("/api/ingresos?quincena=3", headers=h).status_code == 422
    assert api.get("/api/reportes/pdf?quincena=x", headers=h).status_code == 422


def test_sin_token_no_hay_reportes(api):
    for ruta in ("/api/reportes/pdf", "/api/reportes/excel", "/api/reportes/datos", "/api/ingresos"):
        assert api.get(ruta).status_code == 401


# ---------------------------------------------------------------------------------- PDF
def test_pdf_se_genera_en_memoria_con_encabezado_de_periodo(api, auth, crear_usuario, datos):
    r = api.get("/api/reportes/pdf?anio=2026&mes=10&quincena=ambas", headers=auth(crear_usuario("admin.demo", rol="admin")))
    assert r.status_code == 200 and r.content.startswith(b"%PDF") and r.headers["content-type"] == "application/pdf"
    assert 'filename="reporte-2026-10-mes-completo.pdf"' in r.headers["content-disposition"]
    texto = _texto_pdf(r.content)
    if texto is not None:
        for esperado in ("Periodo parcial, corte al 20/10/2026", "1. Resumen general", "2. Por CM", "3. Detalle por paquete",
                         "4. Prórrogas", "5. Pendientes de renovar", "6. Tasa de renovación", "$8,300.00", "VENCIDA"):
            assert esperado in texto, esperado


def test_pdf_periodo_cerrado(api, auth, crear_usuario, datos, hoy_fijo):
    hoy_fijo(dt.date(2026, 11, 3))
    r = api.get("/api/reportes/pdf?anio=2026&mes=10&quincena=1", headers=auth(crear_usuario("admin.demo", rol="admin")))
    texto = _texto_pdf(r.content)
    if texto is not None:
        assert "Periodo cerrado" in texto and "1ra quincena de octubre 2026" in texto.replace("\n", " ")


def _texto_pdf(contenido: bytes):
    if not shutil.which("pdftotext"):
        return None
    out = subprocess.run(["pdftotext", "-layout", "-", "-"], input=contenido, capture_output=True, check=True).stdout
    return out.decode()


def test_admin_solo_lectura_descarga_reportes(api, auth, crear_usuario, datos):
    h = auth(crear_usuario("lectura.demo", rol="admin", solo_lectura=True))
    assert api.get("/api/reportes/pdf?anio=2026&mes=10&quincena=ambas", headers=h).status_code == 200
    assert api.get("/api/reportes/excel?anio=2026&mes=10&quincena=ambas", headers=h).status_code == 200


# -------------------------------------------------------------------------------- Excel
def _excel(api, h, q="ambas"):
    r = api.get(f"/api/reportes/excel?anio=2026&mes=10&quincena={q}", headers=h)
    assert r.status_code == 200
    return r.content


def test_excel_una_hoja_por_seccion_y_formulas_reales_en_totales(api, auth, crear_usuario, datos):
    wb = load_workbook(io.BytesIO(_excel(api, auth(crear_usuario("admin.demo", rol="admin")))))
    assert wb.sheetnames == ["Resumen", "Por CM", "Detalle paquetes", "Prórrogas", "Pendientes", "Tasa renovación"]
    det, cm, pr, ta = wb["Detalle paquetes"], wb["Por CM"], wb["Prórrogas"], wb["Tasa renovación"]
    fila_tot = det.max_row
    assert [det.cell(fila_tot, c).value for c in (5, 6, 7)] == [f"=SUM(E2:E{fila_tot - 1})", f"=SUM(F2:F{fila_tot - 1})",
                                                                f"=SUM(G2:G{fila_tot - 1})"]
    assert det["G2"].value == "=E2-F2"                                           # restante: fórmula, no valor
    assert str(cm.cell(cm.max_row, 4).value).startswith("=SUM(") and str(cm["D2"].value).startswith("=SUMIFS(")
    assert str(cm.cell(cm.max_row, 7).value).startswith("=IF(")
    assert str(pr.cell(pr.max_row, 6).value).startswith("=SUM(") and pr["H2"].value.startswith("=G2-Resumen!")
    assert str(ta.cell(ta.max_row, 4).value).startswith("=IF(")
    res = wb["Resumen"]
    assert res["B7"].value.startswith("='Por CM'!D") and res["B9"].value == "=B7-B8"


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice no disponible para recalcular fórmulas")
def test_excel_las_formulas_calculan_los_mismos_numeros_que_el_json(api, auth, crear_usuario, datos, tmp_path):
    h = auth(crear_usuario("admin.demo", rol="admin"))
    (tmp_path / "in").mkdir()
    (tmp_path / "in" / "r.xlsx").write_bytes(_excel(api, h))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"),
                    str(tmp_path / "in" / "r.xlsx")], check=True, capture_output=True, timeout=120)
    wb = load_workbook(tmp_path / "out" / "r.xlsx", data_only=True)
    res = wb["Resumen"]
    assert (res["B7"].value, res["B8"].value, res["B9"].value) == (8300, 3000, 5300)
    assert round(res["B10"].value, 3) == round(3000 / 8300, 3)
    cm = {r[0].value: [c.value for c in r] for r in wb["Por CM"].iter_rows(min_row=2, max_row=wb["Por CM"].max_row - 1)}
    assert cm["Ana"][2:6] == [3, 6000, 1500, 4500] and cm["Beto"][3] == 1500
    total_cm = [c.value for c in wb["Por CM"][wb["Por CM"].max_row]]
    assert total_cm[3:6] == [8300, 3000, 5300]
    pr = wb["Prórrogas"]
    assert [pr.cell(r, 8).value for r in (2, 3)] == [-1, 5] and pr.cell(pr.max_row, 6).value == 1500 + 400
    assert pr["I2"].value.startswith("VENCIDA")
    ta = wb["Tasa renovación"]
    assert [ta.cell(ta.max_row, c).value for c in (2, 3)] == [4, 1] and round(ta.cell(ta.max_row, 4).value, 3) == 0.25
