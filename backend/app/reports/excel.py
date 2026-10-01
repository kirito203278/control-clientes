"""Reporte Excel: una hoja por sección y FÓRMULAS REALES en los totales (no valores pegados)."""
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.reports.formato import fecha

MORADO, AMARILLO, SOFT, ROJO_BG = "4B2A82", "FFDE00", "F7F2FD", "FCE8E6"
DINERO, PCT, FECHA = '"$"#,##0.00', "0.0%", "dd/mm/yyyy"
ENC = Font(bold=True, color="FFFFFF")
BORDE = Border(bottom=Side(style="thin", color="E3DDEE"))
SEMAFORO = {"verde": "Verde", "amarillo": "Amarillo", "rojo": "Rojo", "gris": "Gris"}
ESTADO = {"activo": "Activo", "por_vencer": "Por vencer", "vencido": "Vencido", "renovado": "Renovado",
          "archivado": "Archivado", "eliminado": "Eliminado"}


def _encabezado(ws, fila, titulos, anchos):
    for i, (t, a) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font, c.fill = ENC, PatternFill("solid", fgColor=MORADO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.row_dimensions[fila].height = 28
    ws.freeze_panes = f"A{fila + 1}"      # (ws.cell() crearía una fila vacía y desplazaría los datos)


def _total(ws, fila, ncols, formulas: dict[int, tuple[str, str]], etiqueta_col=1, etiqueta="Total"):
    for c in range(1, ncols + 1):
        cel = ws.cell(row=fila, column=c)
        cel.fill, cel.font = PatternFill("solid", fgColor=AMARILLO), Font(bold=True)
    ws.cell(row=fila, column=etiqueta_col, value=etiqueta)
    for col, (f, fmt) in formulas.items():
        cel = ws.cell(row=fila, column=col, value=f)
        cel.number_format = fmt


def _cebra(ws, desde, hasta, ncols):
    for r in range(desde, hasta + 1):
        for c in range(1, ncols + 1):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDE
            if (r - desde) % 2:
                cel.fill = PatternFill("solid", fgColor=SOFT)


def generar_excel(datos: dict, generado: dt.datetime) -> bytes:
    per = datos["periodo"]
    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    ws_cm, ws_det = wb.create_sheet("Por CM"), wb.create_sheet("Detalle paquetes")
    ws_pr, ws_pe = wb.create_sheet("Prórrogas"), wb.create_sheet("Pendientes")
    ws_ta = wb.create_sheet("Tasa renovación")

    # ---------------------------------------------------------- Detalle por paquete (base de las fórmulas)
    det = datos["detalle"]
    _encabezado(ws_det, 1, ["CM", "Cliente", "Paquete", "Tipo", "Costo", "Pagado", "Restante", "Fecha de renovación",
                            "Estado", "Semáforo"], [20, 30, 14, 12, 14, 14, 14, 16, 14, 12])
    for i, f in enumerate(det, start=2):
        ws_det.append([f["cm"], f["cliente"], f["paquete"], f["tipo"], float(f["costo"]), float(f["pagado"]), f"=E{i}-F{i}",
                       f["fecha_renovacion"], ESTADO.get(f["estado_efectivo"], f["estado_efectivo"]), SEMAFORO[f["semaforo"]]])
    ult_det = max(len(det) + 1, 2)
    _cebra(ws_det, 2, ult_det, 10)
    for r in range(2, len(det) + 2):
        for col in "EFG":
            ws_det[f"{col}{r}"].number_format = DINERO
        ws_det[f"H{r}"].number_format = FECHA
    fila_tot_det = len(det) + 2
    _total(ws_det, fila_tot_det, 10, {5: (f"=SUM(E2:E{ult_det})", DINERO), 6: (f"=SUM(F2:F{ult_det})", DINERO),
                                      7: (f"=SUM(G2:G{ult_det})", DINERO)}, etiqueta_col=2)
    ws_det.auto_filter.ref = f"A1:J{ult_det}"

    # ----------------------------------------------------------------------------------- Por CM
    cms = datos["por_cm"]
    _encabezado(ws_cm, 1, ["CM", "Clientes", "Paquetes", "Proyección", "Cobrado", "Pendiente", "% cobrado"],
                [26, 12, 12, 16, 16, 16, 13])
    rng = lambda col: f"'Detalle paquetes'!${col}$2:${col}${ult_det}"  # noqa: E731
    for i, c in enumerate(cms, start=2):
        ws_cm.append([c["cm"], c["clientes"], f"=COUNTIFS({rng('A')},A{i})", f"=SUMIFS({rng('E')},{rng('A')},A{i})",
                      f"=SUMIFS({rng('F')},{rng('A')},A{i})", f"=D{i}-E{i}", f"=IF(D{i}=0,0,E{i}/D{i})"])
        for col, fmt in (("D", DINERO), ("E", DINERO), ("F", DINERO), ("G", PCT)):
            ws_cm[f"{col}{i}"].number_format = fmt
    ult_cm = max(len(cms) + 1, 2)
    _cebra(ws_cm, 2, ult_cm, 7)
    fila_tot_cm = len(cms) + 2
    _total(ws_cm, fila_tot_cm, 7, {2: (f"=SUM(B2:B{ult_cm})", "0"), 3: (f"=SUM(C2:C{ult_cm})", "0"),
                                   4: (f"=SUM(D2:D{ult_cm})", DINERO), 5: (f"=SUM(E2:E{ult_cm})", DINERO),
                                   6: (f"=D{fila_tot_cm}-E{fila_tot_cm}", DINERO),
                                   7: (f"=IF(D{fila_tot_cm}=0,0,E{fila_tot_cm}/D{fila_tot_cm})", PCT)})

    # ------------------------------------------------------------------------------- Resumen
    ws_res.column_dimensions["A"].width, ws_res.column_dimensions["B"].width = 30, 34
    ws_res["A1"] = "Reporte de ingresos y renovaciones"
    ws_res["A1"].font = Font(bold=True, size=16, color=MORADO)
    ws_res["A2"], ws_res["A3"] = per["etiqueta"].capitalize(), per["estado_texto"]
    ws_res["A3"].font = Font(bold=True, color="E2530A" if not per["cerrado"] else "1F9D55")
    ws_res["A4"], ws_res["A5"] = f"Alcance: {datos['alcance']}", f"Generado el {generado.strftime('%d/%m/%Y %H:%M')}"
    filas = [("Proyección total", f"='Por CM'!D{fila_tot_cm}", DINERO), ("Cobrado", f"='Por CM'!E{fila_tot_cm}", DINERO),
             ("Pendiente", "=B7-B8", DINERO), ("% cobrado", "=IF(B7=0,0,B8/B7)", PCT),
             ("Fecha de corte", per["corte"], FECHA), ("Estado del periodo", per["estado_texto"], "@"),
             ("Renovaciones del", per["desde"], FECHA), ("al", per["hasta"], FECHA)]
    for i, (k, v, fmt) in enumerate(filas, start=7):
        ws_res.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws_res.cell(row=i, column=2, value=v)
        c.number_format, c.alignment = fmt, Alignment(horizontal="left")
    ws_res["A16"] = ("Definiciones: la proyección es el costo de los paquetes cuya fecha de renovación cae en el periodo; "
                     "cobrado es lo pagado de esos paquetes hasta la fecha de corte; pendiente = proyección - cobrado.")
    ws_res["A16"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_res.merge_cells("A16:B19")

    # ------------------------------------------------------------------------------ Prórrogas
    pr = datos["prorrogas"]
    _encabezado(ws_pr, 1, ["CM", "Cliente", "Paquete", "Costo", "Pagado", "Restante por cobrar", "Fecha límite",
                           "Días (+ faltan / − retraso)", "Situación"], [20, 30, 14, 14, 14, 18, 14, 18, 30])
    for i, r in enumerate(pr, start=2):
        ws_pr.append([r["cm"], r["cliente"], r["paquete"], float(r["costo"]), float(r["pagado"]), f"=D{i}-E{i}",
                      r["fecha_limite"], f"=G{i}-Resumen!$B$11",
                      f'=IF(H{i}<0,"VENCIDA · "&-H{i}&" día(s) de retraso",IF(H{i}=0,"Vence hoy","Faltan "&H{i}&" día(s)"))'])
        for col in "DEF":
            ws_pr[f"{col}{i}"].number_format = DINERO
        ws_pr[f"G{i}"].number_format = FECHA
    ult_pr = max(len(pr) + 1, 2)
    _cebra(ws_pr, 2, ult_pr, 9)
    ws_pr.conditional_formatting.add(f"A2:I{ult_pr}", FormulaRule(formula=["$H2<0"], fill=PatternFill("solid", bgColor=ROJO_BG),
                                                                 font=Font(color="C62828", bold=True)))
    _total(ws_pr, len(pr) + 2, 9, {6: (f"=SUM(F2:F{ult_pr})", DINERO)}, etiqueta_col=2, etiqueta="Total por cobrar")

    # ----------------------------------------------------------------------------- Pendientes
    _encabezado(ws_pe, 1, ["Categoría", "CM", "Cliente", "Paquete", "Restante", "Fecha de renovación", "Semáforo"],
                [24, 22, 30, 14, 14, 18, 12])
    fila = 2
    for clave, nombre in (("por_vencer", "Por vencer"), ("vencidos", "Vencidos sin decisión"),
                          ("renovados_sin_pago", "Renovados sin pago")):
        for g in datos["pendientes"][clave]:
            for f in g["items"]:
                ws_pe.append([nombre, g["cm"], f["cliente"], f["paquete"], float(f["restante"]), f["fecha_renovacion"],
                              SEMAFORO[f["semaforo"]]])
                ws_pe[f"E{fila}"].number_format, ws_pe[f"F{fila}"].number_format = DINERO, FECHA
                fila += 1
    ult_pe = max(fila - 1, 2)
    _cebra(ws_pe, 2, ult_pe, 7)
    _total(ws_pe, fila, 7, {5: (f"=SUM(E2:E{ult_pe})", DINERO)}, etiqueta_col=3)

    # -------------------------------------------------------------------- Tasa de renovación
    tr = datos["tasa_renovacion"]
    _encabezado(ws_ta, 1, ["CM", "Vencidos del periodo", "Renovados", "Tasa de renovación"], [28, 20, 14, 20])
    for i, r in enumerate(tr, start=2):
        ws_ta.append([r["cm"], r["vencidos"], r["renovados"], f"=IF(B{i}=0,0,C{i}/B{i})"])
        ws_ta[f"D{i}"].number_format = PCT
    ult_ta = max(len(tr) + 1, 2)
    _cebra(ws_ta, 2, ult_ta, 4)
    ft = len(tr) + 2
    _total(ws_ta, ft, 4, {2: (f"=SUM(B2:B{ult_ta})", "0"), 3: (f"=SUM(C2:C{ult_ta})", "0"),
                          4: (f"=IF(B{ft}=0,0,C{ft}/B{ft})", PCT)})

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
